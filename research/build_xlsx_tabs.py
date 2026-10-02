"""Workstream tabs for ADBE_alt_data.xlsx. Each add_* function is guarded: missing files -> tab with a note."""
import os, glob
import pandas as pd
from openpyxl.styles import Font

def _num(v):
    try:
        if v is None or (isinstance(v, float) and pd.isna(v)): return None
        return float(v)
    except Exception:
        return None

def note(ws, r, text):
    ws.cell(row=r, column=1, value=text).font = Font(italic=True, color="7F8C8D")

def add_trends_competitors(wb, ctx):
    read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT = ctx
    ws = wb.create_sheet("Trends_Competitors")
    ws["A1"] = "WS5: competitor growth vs Adobe (blue = sourced input; black = formula)"; ws["A1"].font = BOLD
    df = read_csv("ws5_trends_competitors/data/growth_comparison.csv")
    r = 3
    if df is None:
        note(ws, r, "UNAVAILABLE: growth_comparison.csv missing"); return
    note(ws, r, "Source: research/ws5_trends_competitors/data/growth_comparison.csv; each row's primary URL is in the Sources tab (WS5). Google Trends: UNAVAILABLE (host blocked); manual export kit in ws5 data/."); r += 1
    cols = ["entity","metric","period","period_end","value_musd","prior_value_musd","yoy_growth_calc","incremental_calc","annualized_incremental_calc","basis","note"]
    hdr(ws, r, cols); r += 1
    for _, rec in df.iterrows():
        ws.cell(row=r, column=1, value=rec["entity"]); ws.cell(row=r, column=2, value=rec["metric"])
        ws.cell(row=r, column=3, value=rec["period"]); ws.cell(row=r, column=4, value=rec["period_end"])
        v = _num(rec["value_musd"]); p = _num(rec["prior_value_musd"])
        c = ws.cell(row=r, column=5, value=v); c.font = BLUE
        c = ws.cell(row=r, column=6, value=p); c.font = BLUE
        if v is not None and p:
            ws.cell(row=r, column=7, value=f"=E{r}/F{r}-1").number_format = "0.0%"
            ws.cell(row=r, column=8, value=f"=E{r}-F{r}")
            mult = 4 if str(rec["period"]).startswith("Q") else 1
            ws.cell(row=r, column=9, value=f"=H{r}*{mult}")
        elif _num(rec["yoy_growth_pct"]) is not None:
            c = ws.cell(row=r, column=7, value=_num(rec["yoy_growth_pct"])/100); c.font = BLUE; c.number_format = "0.0%"
        ws.cell(row=r, column=10, value=rec["basis"]); ws.cell(row=r, column=11, value=rec["note"])
        r += 1
    # share-of-pool block
    r += 1
    ws.cell(row=r, column=1, value="Share-of-pool calc (latest annualized run-rates; formulas)").font = BOLD; r += 1
    hdr(ws, r, ["entity","run_rate_musd (latest quarter x4 / ARR)","annualized_increment_musd","share_of_stock","share_of_increment"]); r += 1
    rows = {"Adobe C&MP (Q3 FY26 x4)": (4651*4, 2136), "Figma (Q2'26 x4)": (370.1*4, 482.0), "Canva (Q2'26 revenue x4)": (921.9*4, 742.4)}
    first = r
    for k, (rr, inc) in rows.items():
        ws.cell(row=r, column=1, value=k)
        ws.cell(row=r, column=2, value=rr).font = BLUE
        ws.cell(row=r, column=3, value=inc).font = BLUE
        r += 1
    last = r - 1
    for rr_ in range(first, last + 1):
        ws.cell(row=rr_, column=4, value=f"=B{rr_}/SUM($B${first}:$B${last})").number_format = "0.0%"
        ws.cell(row=rr_, column=5, value=f"=C{rr_}/SUM($C${first}:$C${last})").number_format = "0.0%"
    ws.cell(row=r, column=1, value="Pool total"); ws.cell(row=r, column=2, value=f"=SUM(B{first}:B{last})"); ws.cell(row=r, column=3, value=f"=SUM(C{first}:C{last})")
    ws.cell(row=r+1, column=1, value="Pool growth (increment / (stock - increment))"); ws.cell(row=r+1, column=2, value=f"=C{r}/(B{r}-C{r})").number_format = "0.0%"
    autosize(ws)

def add_sources(wb, ctx):
    read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT = ctx
    ws = wb.create_sheet("Sources")
    ws["A1"] = "All workstream sources (concatenated from each ws*/sources.csv)"; ws["A1"].font = BOLD
    hdr(ws, 3, ["workstream","claim","value","url","date_accessed","notes"])
    r = 4
    for p in sorted(glob.glob(os.path.join(ROOT, "ws*_*", "sources.csv"))):
        wsname = os.path.basename(os.path.dirname(p))
        try:
            df = pd.read_csv(p)
        except Exception as e:
            ws.cell(row=r, column=1, value=wsname); ws.cell(row=r, column=2, value=f"ERROR reading sources.csv: {e}"); r += 1; continue
        for _, rec in df.iterrows():
            ws.cell(row=r, column=1, value=wsname)
            for j, col in enumerate(["claim","value","url","date_accessed","notes"]):
                v = rec.get(col)
                ws.cell(row=r, column=2 + j, value=None if (isinstance(v, float) and pd.isna(v)) else (str(v) if v is not None else None))
            r += 1
    autosize(ws, maxw=80)

# placeholders, filled in as workstreams complete
def add_price_volume(wb, ctx):
    read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT = ctx
    ws = wb.create_sheet("PriceVolume")
    ws["A1"] = "WS1: price vs volume decomposition"; ws["A1"].font = BOLD
    note(ws, 3, "PENDING: workstream output not yet integrated")

def add_seat_compression(wb, ctx):
    read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT = ctx
    ws = wb.create_sheet("SeatCompression")
    ws["A1"] = "WS2: revenue per creative worker / seat compression"; ws["A1"].font = BOLD
    note(ws, 3, "PENDING: workstream output not yet integrated")

def add_task_exposure(wb, ctx):
    read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT = ctx
    ws = wb.create_sheet("TaskExposure")
    ws["A1"] = "WS3: AI task exposure by creative-workflow stage"; ws["A1"].font = BOLD
    note(ws, 3, "PENDING: workstream output not yet integrated")

def add_call_nlp(wb, ctx):
    read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT = ctx
    ws = wb.create_sheet("CallNLP")
    ws["A1"] = "WS4: earnings-call analyst-question analysis and forward P/E"; ws["A1"].font = BOLD
    note(ws, 3, "PENDING: workstream output not yet integrated")

def add_tabs(wb, name_row, read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT):
    ctx = (read_csv, write_df, hdr, autosize, BLUE, BLACK, BOLD, ROOT)
    for fn in (add_price_volume, add_seat_compression, add_task_exposure, add_call_nlp, add_trends_competitors, add_sources):
        try:
            fn(wb, ctx)
        except Exception as e:
            ws = wb.create_sheet(fn.__name__.replace("add_", "ERR_")[:28])
            ws["A1"] = f"ERROR building tab: {e}"
            print("TAB ERROR", fn.__name__, e)
